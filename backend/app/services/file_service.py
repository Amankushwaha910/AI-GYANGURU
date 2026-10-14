"""
File service — upload, validation, text extraction, OCR pipeline.
"""
import asyncio
import io
import logging
import math
import mimetypes
from typing import Optional, Tuple
from uuid import UUID

from fastapi import BackgroundTasks, UploadFile
from sqlalchemy import update
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.config import settings
from app.core.exceptions import FileProcessingError, NotFoundError, ValidationError
from app.models.analytics import EventType
from app.models.history import HistoryItemType
from app.models.upload import Upload, UploadStatus
from app.ocr.engine import get_ocr_engine
from app.repositories.analytics_repository import AnalyticsRepository
from app.repositories.history_repository import HistoryRepository
from app.repositories.upload_repository import UploadRepository
from app.storage.supabase_storage import get_storage_client

logger = logging.getLogger("gyanguru.file_service")

# Maximum characters sent to AI to avoid token-limit issues.
# Roughly 6 000 words / ~8 000 tokens — well within every supported model.
MAX_EXTRACTED_CHARS = 24_000

# Scanned PDFs are rendered one page at a time. Cap each RGB raster at roughly
# 36 MiB and cap concurrent PDF OCR jobs so background uploads cannot create
# an unbounded set of large page buffers.
MAX_OCR_PAGE_PIXELS = 12_000_000
MAX_CONCURRENT_PDF_OCR = 2
_PDF_OCR_SEMAPHORE = asyncio.Semaphore(MAX_CONCURRENT_PDF_OCR)


class FileService:
    def __init__(self, db: AsyncSession) -> None:
        self.db = db
        self.upload_repo = UploadRepository(db)
        self.analytics_repo = AnalyticsRepository(db)
        self.history_repo = HistoryRepository(db)
        self.storage = get_storage_client()
        self.ocr_engine = get_ocr_engine()

    # ── Public API ─────────────────────────────────────────────────────────────

    async def upload_and_process(
        self,
        user_id: UUID,
        file: UploadFile,
        background_tasks: BackgroundTasks,
    ) -> Upload:
        """
        Upload pipeline:
        1. Validate file (type, size)
        2. Read bytes
        3. Store to Supabase Storage
        4. Create Upload DB record (status=pending)
        5. Queue background text extraction
        6. Log analytics + history
        """
        content = await self._validate_file(file)
        filename = file.filename or "upload"
        ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
        mime_type = (
            file.content_type
            or mimetypes.guess_type(filename)[0]
            or "application/octet-stream"
        )

        # Authentication has already read the user in this request. Release
        # that read transaction before waiting on the remote storage service.
        await self.db.commit()
        storage_path = await self.storage.upload(
            bucket=settings.supabase_storage_bucket,
            user_id=str(user_id),
            filename=filename,
            content=content,
            mime_type=mime_type,
        )
        try:
            upload = await self.upload_repo.create(
                user_id=user_id,
                original_filename=filename,
                storage_path=storage_path,
                file_type=ext,
                mime_type=mime_type,
                file_size_bytes=len(content),
                status=UploadStatus.pending,
            )

            # Queue extraction only after DB and history work succeeds.
            # Pass serialisable primitives, never ORM objects or DB sessions.
            await self.analytics_repo.log_event(
                user_id=user_id,
                event_type=EventType.file_uploaded,
                topic=filename,
                module="file",
                resource_id=str(upload.id),
            )
            await self.history_repo.add_entry(
                user_id=user_id,
                item_type=HistoryItemType.upload,
                topic=filename,
                resource_id=str(upload.id),
            )
            background_tasks.add_task(
                _extract_text_background,
                upload_id=upload.id,
                content=content,
                file_type=ext,
                mime_type=mime_type,
            )
            return upload
        except BaseException:
            # DB writes are still in the request transaction. Roll them back
            # and compensate for the already-created remote object.
            try:
                await self.db.rollback()
            except Exception:
                logger.error("file_service: DB rollback failed after upload persistence failure")
            try:
                await self.storage.delete(
                    bucket=settings.supabase_storage_bucket,
                    path=storage_path,
                )
            except Exception:
                logger.error(
                    "file_service: storage cleanup failed after upload persistence failure"
                )
            raise

    async def get_upload(self, user_id: UUID, upload_id: UUID) -> Upload:
        upload = await self.upload_repo.get_user_upload(user_id, upload_id)
        if not upload:
            raise NotFoundError("Upload", str(upload_id))
        return upload

    async def list_uploads(self, user_id: UUID, page: int, page_size: int):
        return await self.upload_repo.get_user_uploads(
            user_id=user_id, page=page, page_size=page_size
        )

    async def delete_upload(self, user_id: UUID, upload_id: UUID) -> None:
        upload = await self.upload_repo.get_user_upload(user_id, upload_id)
        if not upload:
            raise NotFoundError("Upload", str(upload_id))
        # Release the connection before making the remote storage request.
        await self.db.commit()
        try:
            await self.storage.delete(
                bucket=settings.supabase_storage_bucket,
                path=upload.storage_path,
            )
        except Exception as exc:
            # Keep the DB row so the owner can retry; deleting it here would
            # orphan a still-live object and make its storage path unrecoverable.
            logger.warning("file_service: remote delete failed; retaining upload row")
            raise FileProcessingError(
                "The file could not be deleted from storage. Please try again."
            ) from exc
        await self.upload_repo.delete(upload)

    # ── Validation ─────────────────────────────────────────────────────────────

    async def _validate_file(self, file: UploadFile) -> bytes:
        if not file.filename:
            raise ValidationError("File must have a filename")

        ext = file.filename.rsplit(".", 1)[-1].lower() if "." in file.filename else ""
        if ext not in settings.allowed_extensions_list:
            raise ValidationError(
                f"File type '.{ext}' is not supported. "
                f"Allowed: {', '.join(settings.allowed_extensions_list)}"
            )

        max_size = settings.max_file_size_bytes
        declared_size = getattr(file, "size", None)
        if declared_size is not None and declared_size > max_size:
            raise ValidationError(
                f"File size exceeds maximum of {settings.max_file_size_mb} MB"
            )

        # Bound the read even when UploadFile.size is unavailable or untrusted.
        content = await file.read(max_size + 1)
        if len(content) > max_size:
            raise ValidationError(
                f"File size exceeds maximum of {settings.max_file_size_mb} MB"
            )
        await file.seek(0)
        return content


# ── Background task (module-level, not a bound method) ────────────────────────
#
# Defined at module level (not as FileService method) so:
#   • no reference to the request-scoped FileService instance or its DB session
#   • receives only plain Python primitives (UUID, bytes, str) — fully picklable
#   • opens its own DB session via get_session_factory()
#   • uses direct UPDATE statements (not ORM flush/refresh) to avoid the
#     PgBouncer NullPool transaction bug that caused the "stuck at Pending" issue
#
async def _extract_text_background(
    upload_id: UUID,
    content: bytes,
    file_type: str,
    mime_type: str,
) -> None:
    """
    Background task: extract text from file bytes and update the Upload record.

    Uses raw SQL UPDATE statements so that each DB write is a self-contained
    transaction — compatible with PgBouncer's NullPool / transaction-pooler mode.
    """
    from app.core.database import get_session_factory

    factory = get_session_factory()

    # ── Step 1: mark as processing ──────────────────────────────────────────
    async with factory() as session:
        try:
            await session.execute(
                update(Upload)
                .where(Upload.id == upload_id)
                .values(status=UploadStatus.processing)
            )
            await session.commit()
        except Exception as exc:
            await session.rollback()
            logger.error(
                "file_service: failed to set status=processing for upload %s: %s",
                upload_id, exc,
            )
            return

    # ── Step 2: extract text ────────────────────────────────────────────────
    try:
        extracted_text, ocr_used = await _extract_text(content, file_type, mime_type)
        # Truncate to safe AI context size
        if len(extracted_text) > MAX_EXTRACTED_CHARS:
            extracted_text = extracted_text[:MAX_EXTRACTED_CHARS]
            logger.info(
                "file_service: extracted text truncated to %d chars for upload %s",
                MAX_EXTRACTED_CHARS, upload_id,
            )
        error_message = None
    except Exception as exc:
        logger.error(
            "file_service: text extraction failed for upload %s: %s",
            upload_id, exc,
        )
        extracted_text = None
        ocr_used = False
        error_message = str(exc)

    # ── Step 3: persist result ──────────────────────────────────────────────
    if extracted_text is not None:
        new_status = UploadStatus.completed
        values = {
            "status": new_status,
            "extracted_text": extracted_text,
            "ocr_used": ocr_used,
            "error_message": None,
        }
    else:
        new_status = UploadStatus.failed
        values = {
            "status": new_status,
            "error_message": error_message or "Text extraction failed",
        }

    async with factory() as session:
        try:
            await session.execute(
                update(Upload)
                .where(Upload.id == upload_id)
                .values(**values)
            )
            await session.commit()
            logger.info(
                "file_service: upload %s → %s (ocr=%s)",
                upload_id, new_status, ocr_used,
            )
        except Exception as exc:
            await session.rollback()
            logger.error(
                "file_service: failed to persist extraction result for upload %s: %s",
                upload_id, exc,
            )


# ── Text extraction helpers (pure functions, no DB) ───────────────────────────

async def _extract_text(
    content: bytes,
    file_type: str,
    mime_type: str,
) -> Tuple[str, bool]:
    """
    Extract text from file bytes.
    Returns (extracted_text, ocr_used).
    Raises FileProcessingError on failure.
    """
    ocr_used = False

    if file_type == "pdf":
        # PyMuPDF parsing is synchronous and may be CPU-heavy for large files.
        text = await asyncio.to_thread(_extract_from_pdf, content)
        if len(text.strip()) < 50:
            # PDF has no embedded text → try OCR on each page as image
            text = await _ocr_pdf_pages(content)
            ocr_used = True

    elif file_type == "docx":
        # python-docx parsing is synchronous; keep it off the event loop.
        text = await asyncio.to_thread(_extract_from_docx, content)

    elif file_type == "txt":
        text = content.decode("utf-8", errors="ignore")

    elif file_type in ("png", "jpg", "jpeg", "webp"):
        ocr_engine = get_ocr_engine()
        text = await ocr_engine.extract_from_image(content)
        ocr_used = True

    else:
        raise FileProcessingError(f"Unsupported file type: {file_type}")

    stripped = text.strip()
    if not stripped:
        raise FileProcessingError(
            f"No text could be extracted from the {file_type.upper()} file. "
            f"The file may be empty, image-only, or encrypted."
        )

    return stripped, ocr_used


def _extract_from_pdf(content: bytes) -> str:
    """Extract embedded text from PDF using PyMuPDF (fitz)."""
    try:
        import fitz  # PyMuPDF
        doc = fitz.open(stream=content, filetype="pdf")
        parts = []
        for page in doc:
            parts.append(page.get_text())
        doc.close()
        return "\n".join(parts)
    except Exception as exc:
        raise FileProcessingError(f"PDF text extraction failed: {exc}") from exc


async def _ocr_pdf_pages(content: bytes) -> str:
    """
    Render each PDF page as an image and run OCR on it.
    Used when embedded text is absent (scanned/image PDFs).
    """
    try:
        page_count = await asyncio.to_thread(_get_pdf_page_count, content)
        ocr_engine = get_ocr_engine()
        parts = []
        for page_number in range(page_count):
            # Open/render/close the PDF wholly inside one worker call. PyMuPDF
            # objects therefore never cross thread boundaries. Processing one
            # page at a time bounds per-document memory; the semaphore caps
            # concurrent page buffers across background uploads.
            async with _PDF_OCR_SEMAPHORE:
                img_bytes = await _await_task_to_finish(
                    asyncio.create_task(
                        asyncio.to_thread(_render_pdf_page, content, page_number)
                    )
                )
                page_text = await _await_task_to_finish(
                    asyncio.create_task(ocr_engine.extract_from_image(img_bytes))
                )
            if page_text.strip():
                parts.append(page_text.strip())
        return "\n".join(parts)
    except Exception as exc:
        raise FileProcessingError(f"PDF OCR failed: {exc}") from exc


async def _await_task_to_finish(task: "asyncio.Task"):
    """Hold bounded OCR resources until thread-backed work really finishes."""
    try:
        return await asyncio.shield(task)
    except asyncio.CancelledError:
        # Cancelling asyncio.to_thread/run_in_executor does not stop its worker.
        # Keep the semaphore slot until that worker releases its page buffer.
        try:
            await asyncio.shield(task)
        except (asyncio.CancelledError, Exception):
            pass
        raise


def _get_pdf_page_count(content: bytes) -> int:
    """Read page count off-loop without retaining a PyMuPDF document."""
    import fitz

    doc = fitz.open(stream=content, filetype="pdf")
    try:
        return len(doc)
    finally:
        doc.close()


def _pdf_render_zoom(width: float, height: float) -> float:
    """Use the normal 2x OCR scale unless it would exceed the pixel budget."""
    area = max(width, 0) * max(height, 0)
    if not area:
        return 2.0
    zoom = min(2.0, math.sqrt(MAX_OCR_PAGE_PIXELS / area))
    # PyMuPDF rounds raster dimensions up, so account for that rounding too.
    while math.ceil(width * zoom) * math.ceil(height * zoom) > MAX_OCR_PAGE_PIXELS:
        zoom *= 0.999
    return zoom


def _render_pdf_page(content: bytes, page_number: int) -> bytes:
    """Render one bounded-size page to PNG entirely on its worker thread."""
    import fitz

    doc = fitz.open(stream=content, filetype="pdf")
    try:
        page = doc.load_page(page_number)
        zoom = _pdf_render_zoom(page.rect.width, page.rect.height)
        pix = page.get_pixmap(
            matrix=fitz.Matrix(zoom, zoom),
            alpha=False,
        )
        return pix.tobytes("png")
    finally:
        doc.close()


def _extract_from_docx(content: bytes) -> str:
    """Extract paragraph text from DOCX using python-docx."""
    try:
        import docx
        doc = docx.Document(io.BytesIO(content))
        lines = [para.text for para in doc.paragraphs if para.text.strip()]
        # Also extract table cells
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        lines.append(cell.text.strip())
        return "\n".join(lines)
    except Exception as exc:
        raise FileProcessingError(f"DOCX text extraction failed: {exc}") from exc
